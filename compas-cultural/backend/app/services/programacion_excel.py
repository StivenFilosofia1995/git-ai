"""
programacion_excel.py — Importa la programación mensual que llega en Excel
(Fundación Grupo EPM: Parque de los Deseos, Museo del Agua, UVA, Biblioteca EPM).

100 % determinista. Detecta el formato de cada hoja por sus encabezados:
- PARQUE:     "Fecha (DD/MM/AAAA)" + "Abierto al público" (una hoja por mes)
- MUSEO:      "Oferta Abierta/Cerrada"
- UVA:        "UVA" + "Días del mes"
- BIBLIOTECA: "Título del curso" + "Fecha(s)" (el mes viene en el título de la hoja)

Reglas:
- Lo cerrado (grupos ya conformados), lo "pendiente" y lo que no es abierto al
  público NO se publica. Lo "tentativo" entra en cuarentena (oculto) para confirmar.
- Talleres, clubes, eventos y cine: un evento por sesión (aparecen el día real).
- Cursos y semilleros (proceso con inscripción): un solo evento en la fecha de inicio,
  con las sesiones en la descripción (no llenan "Hoy" todos los días).
- Horas imposibles se corrigen sin adivinar: "9:00 p.m. a 12:00 m." → 9:00 a.m.;
  "23:30 a 13:00" → 11:30. Si no se puede, queda "hora por confirmar".
- Ubicación: sedes de seeds/data/sedes_epm.json (solo coordenadas verificadas).
"""
from __future__ import annotations

import io
import json
import re
import struct
import unicodedata
import zipfile
import zlib
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Optional
from zoneinfo import ZoneInfo

CO_TZ = ZoneInfo("America/Bogota")
SEDES_JSON = Path(__file__).resolve().parents[2] / "seeds" / "data" / "sedes_epm.json"

_MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
          "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
_DIAS = {"lunes": 0, "martes": 1, "miercoles": 2, "jueves": 3, "viernes": 4, "sabado": 5, "domingo": 6}
_SERIE = {"curso", "semillero", "diplomado"}  # proceso con inscripción → un solo evento


def _norm(s: Any) -> str:
    s = unicodedata.normalize("NFD", str(s or "").lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def _limpio(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("\xa0", " ")).strip()


# ─── Lectura tolerante del archivo ─────────────────────────────────────────

def _reparar_xlsx(data: bytes) -> bytes:
    """Reconstruye un .xlsx al que le falta el índice final (descarga cortada)."""
    partes: dict[str, bytes] = {}
    i = 0
    while True:
        i = data.find(b"PK\x03\x04", i)
        if i < 0:
            break
        (_v, _f, metodo, _t, _d, _crc, csize, _us, nlen, xlen) = struct.unpack("<HHHHHIIIHH", data[i + 4:i + 30])
        nombre = data[i + 30:i + 30 + nlen].decode("utf-8", "ignore")
        ini = i + 30 + nlen + xlen
        crudo = data[ini:ini + csize] if csize else data[ini:]
        try:
            partes[nombre] = zlib.decompressobj(-15).decompress(crudo) if metodo == 8 else crudo
        except Exception:
            pass
        i = ini + max(csize, 1)
    texto = {k: v.decode("utf-8", "ignore") for k, v in partes.items()}
    # Quitar referencias a partes que no se recuperaron
    for ref in re.findall(r'PartName="/([^"]+)"', texto.get("[Content_Types].xml", "")):
        if ref not in texto:
            texto["[Content_Types].xml"] = re.sub(rf'<Override PartName="/{re.escape(ref)}"[^>]*/>', "",
                                                  texto["[Content_Types].xml"])
    for k in [k for k in texto if k.endswith(".rels")]:
        base = k.replace("_rels/", "").replace(".rels", "")
        carpeta = base.rsplit("/", 1)[0] if "/" in base else ""

        def _existe(target: str) -> bool:
            t = target.lstrip("/")
            if target.startswith("/"):
                return t in texto
            ruta = (carpeta + "/" + t) if carpeta else t
            partes_ruta: list[str] = []
            for seg in ruta.split("/"):
                if seg == "..":
                    partes_ruta and partes_ruta.pop()
                else:
                    partes_ruta.append(seg)
            return "/".join(partes_ruta) in texto

        texto[k] = re.sub(r'<Relationship [^>]*Target="([^"]+)"[^>]*/>',
                          lambda m: m.group(0) if (_existe(m.group(1)) or "http" in m.group(1)) else "", texto[k])
    for k, v in list(texto.items()):
        if k.startswith("xl/worksheets/sheet") and k.endswith(".xml"):
            if not v.rstrip().endswith("</worksheet>"):  # hoja cortada: se conservan las filas completas
                corte = v.rfind("</row>")
                v = v[:corte + 6] + "</sheetData></worksheet>" if corte > 0 else v
            texto[k] = re.sub(r"<legacyDrawing [^>]*/>", "", v)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in texto.items():
            z.writestr(k, v.encode("utf-8"))
    return buf.getvalue()


def abrir_libro(data: bytes):
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True), False
    except (zipfile.BadZipFile, KeyError, Exception):
        return openpyxl.load_workbook(io.BytesIO(_reparar_xlsx(data)), data_only=True, read_only=True), True


# ─── Fechas y horas ─────────────────────────────────────────────────────────

def parsear_hora(v: Any) -> Optional[tuple[int, int]]:
    """time/datetime, '14:00:00', '9:00 a.m', '2:00 p.m', '12:00:00 m', '11:00m', '10:00:009'."""
    if v is None or v == "":
        return None
    if isinstance(v, (time, datetime)):
        return v.hour, v.minute
    t = _norm(v).replace(".", "")
    m = re.search(r"(\d{1,2})\s*:\s*(\d{2})", t)
    if not m:
        m2 = re.search(r"\b(\d{1,2})\s*(am|pm|a m|p m|m)\b", t)
        if not m2:
            return None
        h, mi = int(m2.group(1)), 0
        suf = m2.group(2).replace(" ", "")
    else:
        h, mi = int(m.group(1)), int(m.group(2))
        resto = t[m.end():]
        suf = "pm" if re.search(r"^\s*(:\d+)?\s*p\s?m", resto) else "am" if re.search(r"^\s*(:\d+)?\s*a\s?m", resto) \
            else "m" if re.search(r"^\s*(:\d+\d*)?\s*m\b", resto) else ""
    if suf == "pm" and h < 12:
        h += 12
    elif suf == "am" and h == 12:
        h = 0
    elif suf == "m" and h == 12:
        h = 12  # "12:00 m." = mediodía
    if not (0 <= h <= 23 and 0 <= mi <= 59):
        return None
    return h, mi


def corregir_rango(ini: Optional[tuple[int, int]], fin: Optional[tuple[int, int]]):
    """Si el inicio queda después del fin, el error típico es 12 horas de más."""
    if ini and fin and ini > fin:
        if ini[0] >= 12 and (ini[0] - 12, ini[1]) < fin:
            return (ini[0] - 12, ini[1]), fin
        return None, fin  # no se puede saber: hora por confirmar
    return ini, fin


def _dias_semana_en_mes(anio: int, mes: int, wds: set[int]) -> list[date]:
    d, out = date(anio, mes, 1), []
    while d.month == mes:
        if d.weekday() in wds:
            out.append(d)
        d += timedelta(days=1)
    return out


def parsear_fechas(v: Any, anio: int, mes: Optional[int], dias_txt: str = "") -> tuple[list[date], bool]:
    """Devuelve (fechas, tentativo). Formatos vistos en los Excel de EPM:
    datetime · '06/10/2026 13/10/2026' · '03,10,17y 24/02/2024' · '6,13,20' (mes de la hoja) ·
    '10 y 24 de octubre' · 'Todos los jueves de octubre' (+ columna Día(s)) · 'pendiente' · 'tentativo 18 de octubre'.
    """
    if isinstance(v, datetime):
        return [v.date()], False
    if isinstance(v, date):
        return [v], False
    t = _norm(v)
    if not t or "pendiente" in t or "por definir" in t:
        return [], False
    tentativo = "tentativ" in t or "por confirmar" in t
    # dd/mm/aaaa completos
    completas = re.findall(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if completas:
        out = []
        for d, m, y in completas:
            try:
                out.append(date(int(y), int(m), int(d)))
            except ValueError:
                pass
        # "03,10,17y 24/02/2024": días sueltos antes del último dd/mm/aaaa
        cola = re.match(r"^\s*([\d,\sy]+?)\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$", t)
        if cola and len(completas) == 1:
            m_, y_ = int(cola.group(3)), int(cola.group(4))
            for d in re.findall(r"\d{1,2}", cola.group(1)):
                try:
                    out.append(date(y_, m_, int(d)))
                except ValueError:
                    pass
        return sorted(set(out)), tentativo
    mes_txt = next((n for k, n in _MESES.items() if re.search(rf"\b{k}\b", t)), None)
    mes_ef = mes_txt or mes
    # "Todos los jueves de octubre" / "Todos los sábados" (+ columna Día(s))
    if "todos los" in t or "todas las" in t:
        wds = {n for k, n in _DIAS.items() if re.search(rf"\b{k}s?\b", t + " " + _norm(dias_txt))}
        if wds and mes_ef:
            return _dias_semana_en_mes(anio, mes_ef, wds), tentativo
    nums = [int(x) for x in re.findall(r"\b(\d{1,2})\b", t)]
    if nums and mes_ef:
        out = []
        for d in nums:
            try:
                out.append(date(anio, mes_ef, d))
            except ValueError:
                pass
        return sorted(set(out)), tentativo
    return [], tentativo


# ─── Sedes ──────────────────────────────────────────────────────────────────

def cargar_sedes() -> list[dict]:
    try:
        return json.loads(SEDES_JSON.read_text(encoding="utf-8"))["sedes"]
    except Exception:
        return []


def resolver_sede(texto: str, sedes: list[dict], por_defecto: Optional[str] = None) -> Optional[dict]:
    t = _norm(texto)
    for s in sedes:
        if _norm(s["nombre"]) == t:
            return s
    # alias más largo primero ("uva de la imaginacion" antes que "biblioteca")
    for s, a in sorted(((s, a) for s in sedes for a in s.get("alias", [])), key=lambda x: -len(x[1])):
        if a and a in t:
            return s
    if por_defecto:
        return next((s for s in sedes if s["nombre"] == por_defecto), None)
    return None


# ─── Categorías ─────────────────────────────────────────────────────────────

def categoria(titulo: str, tipo: str = "") -> str:
    tp = _norm(tipo)
    # El tipo formativo manda: "Iniciación musical" (Curso) es un taller, no un concierto
    if any(k in tp for k in ("curso", "club", "semillero", "taller", "clase")):
        return "taller" if "cine" not in _norm(titulo) else "cine"
    t = _norm(f"{tipo} {titulo}")
    reglas = [("cine", ("cine", "proyeccion", "pelicula")), ("musica_en_vivo", ("retreta", "concierto", "musica", "banda", "orquesta")),
              ("danza", ("danza", "baile", "aerobic")), ("poesia", ("poesia",)), ("libreria", ("lectura", "cuento", "letras", "libro", "literari")),
              ("festival", ("fiesta", "festi", "feria", "festival")), ("galeria", ("exposicion", "muestra")),
              ("conferencia", ("conversatorio", "charla", "tertulia")), ("taller", ("taller", "curso", "laboratorio", "club", "semillero", "clase"))]
    for cat, palabras in reglas:
        if any(p in t for p in palabras):
            return cat
    return "taller"


# ─── Evento ─────────────────────────────────────────────────────────────────

@dataclass
class Fila:
    titulo: str
    descripcion: str
    fechas: list[date]
    hora_ini: Optional[tuple[int, int]]
    hora_fin: Optional[tuple[int, int]]
    sede: Optional[dict]
    publico: str = ""
    tipo: str = ""
    lugar_interno: str = ""
    enlace: str = ""
    requiere_inscripcion: bool = False
    tentativo: bool = False
    hoja: str = ""
    fila: int = 0


@dataclass
class Resultado:
    eventos: list[dict] = field(default_factory=list)
    descartados: list[dict] = field(default_factory=list)
    formatos: dict = field(default_factory=dict)
    reparado: bool = False


def _titulo(t: Any) -> str:
    partes = [p.strip(" -–:") for p in str(t or "").replace("\xa0", " ").split("\n") if p.strip(" -–:")]
    if not partes:
        return ""
    if len(partes) >= 2 and len(partes[0]) <= 26:  # "Museo enseña\nKahoot: tras las huellas…"
        return f"{partes[0]}: {' '.join(partes[1:])}"[:140]
    return " ".join(partes)[:140]


def _slug(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:110]


def _a_eventos(f: Fila, hoy: date) -> list[dict]:
    futuras = [d for d in f.fechas if d >= hoy]
    if not futuras:
        return []
    serie = _norm(f.tipo) in _SERIE or (f.requiere_inscripcion and len(f.fechas) > 1 and "taller" not in _norm(f.tipo))
    hora_ok = f.hora_ini is not None
    h, mi = f.hora_ini or (0, 0)

    def fmt_h(hm):
        if not hm:
            return ""
        hh, mm = hm
        return f"{hh % 12 or 12}:{mm:02d} {'a. m.' if hh < 12 else 'p. m.'}"

    horario = f"{fmt_h(f.hora_ini)} a {fmt_h(f.hora_fin)}" if f.hora_ini and f.hora_fin else fmt_h(f.hora_ini)
    extra = [x for x in (
        f"Lugar: {f.lugar_interno}" if f.lugar_interno and f.sede and _norm(f.lugar_interno) != _norm(f.sede["nombre"]) else "",
        f"Público: {f.publico}" if f.publico else "",
        f"Horario: {horario}" if horario else "",
        "Requiere inscripción" + (f": {f.enlace}" if f.enlace else "") if f.requiere_inscripcion else "Entrada libre",
    ) if x]
    base_desc = _limpio(f.descripcion)
    sede = f.sede or {}
    comunes = {
        "categoria_principal": categoria(f.titulo, f.tipo),
        "es_gratuito": True,
        "precio": "Gratis" + (" · con inscripción" if f.requiere_inscripcion else ""),
        "nombre_lugar": sede.get("nombre"),
        "municipio": sede.get("municipio", "medellin"),
        "barrio": sede.get("barrio"),
        "lat": sede.get("lat"),
        "lng": sede.get("lng"),
        "fuente": "fundacion_epm_excel",
        "fuente_url": f.enlace or sede.get("url") or "https://www.grupo-epm.com/site/fundacionepm/",
        "hora_confirmada": hora_ok,
        "verificado": True,
    }
    comunes["categorias"] = [comunes["categoria_principal"]]
    if f.tentativo:
        comunes["oculto"] = True
        comunes["oculto_motivo"] = "excel:fecha_tentativa"

    def evento(d: date, fin_d: Optional[date] = None, desc_extra: str = "") -> dict:
        ini_dt = datetime(d.year, d.month, d.day, h, mi, tzinfo=CO_TZ)
        fin_dt = None
        if f.hora_fin and hora_ok and not fin_d:
            fh, fm = f.hora_fin
            fin_dt = datetime(d.year, d.month, d.day, fh, fm, tzinfo=CO_TZ)
        desc = " · ".join(x for x in (base_desc, desc_extra, *extra) if x)[:1500]
        return {**comunes, "titulo": f.titulo, "descripcion": desc, "fecha_inicio": ini_dt.isoformat(),
                "fecha_fin": fin_dt.isoformat() if fin_dt and fin_dt > ini_dt else None,
                "es_recurrente": len(f.fechas) > 1,
                "slug": _slug(f"{f.titulo}-{sede.get('nombre', '')}-{d.isoformat()}"),
                "_fila": f"{f.hoja}#{f.fila}"}

    if serie:
        sesiones = ", ".join(str(d.day) for d in f.fechas)
        mes = next(k for k, v in _MESES.items() if v == f.fechas[0].month)
        return [evento(futuras[0], desc_extra=f"{len(f.fechas)} sesiones: {sesiones} de {mes}")]
    return [evento(d) for d in futuras]


# ─── Formatos ───────────────────────────────────────────────────────────────

def _filas(ws, desde: int):
    """Filas desde `desde`, deteniéndose tras 30 vacías seguidas: algunas hojas declaran
    más de un millón de filas (la de UVA) y recorrerlas todas tarda minutos."""
    vacias = 0
    for n, r in enumerate(ws.iter_rows(min_row=desde, values_only=True), start=desde):
        if all(x in (None, "") or (isinstance(x, str) and not x.strip()) for x in r):
            vacias += 1
            if vacias >= 30:
                return
            continue
        vacias = 0
        yield n, r


def _encabezado(ws, buscar: list[str], max_filas: int = 6) -> Optional[tuple[int, dict[str, int]]]:
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_filas, values_only=True), start=1):
        cols = {_norm(c): i for i, c in enumerate(row) if c is not None}
        if all(any(b in k for k in cols) for b in buscar):
            return idx, cols
    return None


def _col(cols: dict[str, int], *claves: str) -> Optional[int]:
    for c in claves:
        for k, i in cols.items():
            if c in k:
                return i
    return None


def _mes_de_titulo(texto: str) -> tuple[Optional[int], Optional[int]]:
    t = _norm(texto)
    mes = next((n for k, n in _MESES.items() if re.search(rf"\b{k}\b", t)), None)
    y = re.search(r"\b(20\d{2})\b", t)
    return mes, int(y.group(1)) if y else None


def leer_programacion(data: bytes, nombre_archivo: str = "", hoy: Optional[date] = None) -> Resultado:
    hoy = hoy or datetime.now(CO_TZ).date()
    wb, reparado = abrir_libro(data)
    res = Resultado(reparado=reparado)
    sedes = cargar_sedes()
    nombre_n = _norm(nombre_archivo)

    for ws in wb.worksheets:
        filas: list[Fila] = []
        enc = _encabezado(ws, ["fecha (dd/mm/aaaa)", "abierto al publico"])
        if enc:  # ── PARQUE (una hoja por mes)
            res.formatos[ws.title] = "parque"
            fila_h, c = enc
            mes, anio = _mes_de_titulo(ws.title)
            anio = anio or hoy.year
            if mes and date(anio, mes, 28) < hoy - timedelta(days=31):
                continue  # meses ya pasados del histórico
            for n, r in _filas(ws, fila_h + 1):
                nombre = r[_col(c, "nombre de la actividad")]
                if not nombre:
                    continue
                abierto = _norm(r[_col(c, "abierto al publico")])
                fechas, tent = parsear_fechas(r[_col(c, "fecha (dd/mm")], anio, mes, str(r[_col(c, "recurrencia")] or ""))
                if abierto.startswith("no"):
                    res.descartados.append({"fila": f"{ws.title}#{n}", "titulo": _titulo(nombre), "motivo": "no_abierto_al_publico"})
                    continue
                if not fechas:
                    res.descartados.append({"fila": f"{ws.title}#{n}", "titulo": _titulo(nombre), "motivo": "fecha_pendiente"})
                    continue
                hi, hf = corregir_rango(parsear_hora(r[_col(c, "hora inicio")]), parsear_hora(r[_col(c, "hora fin")]))
                espacio = _limpio(r[_col(c, "espacio")])
                filas.append(Fila(_titulo(nombre), r[_col(c, "descripcion")] or "", fechas, hi, hf,
                                  resolver_sede("parque de los deseos", sedes), _limpio(r[_col(c, "tipo de publico")]),
                                  _limpio(r[_col(c, "tipo")]), espacio, tentativo=tent, hoja=ws.title, fila=n))
        elif (enc := _encabezado(ws, ["oferta abierta/cerrada"])):  # ── MUSEO DEL AGUA
            res.formatos[ws.title] = "museo"
            fila_h, c = enc
            prev_t, prev_d = "", ""
            for n, r in _filas(ws, fila_h + 1):
                if all(x in (None, "") for x in r):
                    continue
                titulo = _titulo(r[_col(c, "nombre de la actividad")]) or prev_t
                desc = _limpio(r[_col(c, "descripcion")]) or prev_d
                prev_t, prev_d = titulo, desc
                if not titulo:
                    continue
                if _norm(r[_col(c, "oferta abierta")]).startswith("cerrad"):
                    res.descartados.append({"fila": f"{ws.title}#{n}", "titulo": titulo, "motivo": "grupo_cerrado"})
                    continue
                ini_v, fin_v = r[_col(c, "inicio de la actividad")], r[_col(c, "fin de la actividad")]
                fechas = [ini_v.date()] if isinstance(ini_v, datetime) else parsear_fechas(ini_v, hoy.year, None)[0]
                if isinstance(ini_v, datetime) and isinstance(fin_v, datetime) and fin_v.date() > ini_v.date():
                    d = ini_v.date()
                    while d <= fin_v.date():
                        fechas.append(d) if d not in fechas else None
                        d += timedelta(days=1)
                hi, hf = corregir_rango(parsear_hora(r[_col(c, "hora de inicio")]), parsear_hora(r[_col(c, "hora de finalizacion")]))
                insc = _norm(r[_col(c, "requiere inscripcion")]).startswith("si")
                enlace = _limpio(r[_col(c, "enlace de inscripcion")])
                filas.append(Fila(titulo, desc, fechas, hi, hf, resolver_sede("museo del agua", sedes),
                                  _limpio(r[_col(c, "tipo de publico")]), "", "",
                                  enlace if enlace.startswith("http") else "", insc, hoja=ws.title, fila=n))
        elif (enc := _encabezado(ws, ["uva", "dias del mes"])):  # ── UVA
            res.formatos[ws.title] = "uva"
            fila_h, c = enc
            for n, r in _filas(ws, fila_h + 1):
                nombre = r[_col(c, "nombre de la actividad")]
                if not nombre:
                    continue
                titulo = _titulo(nombre)
                if _norm(r[_col(c, "abierto/cerrado")]).startswith("cerrad"):
                    res.descartados.append({"fila": f"{ws.title}#{n}", "titulo": titulo, "motivo": "grupo_cerrado"})
                    continue
                f_ini = r[_col(c, "fecha inicio")]
                f_fin = r[_col(c, "fecha fin")]
                anio = f_ini.year if isinstance(f_ini, datetime) else hoy.year
                mes = f_ini.month if isinstance(f_ini, datetime) else None
                fechas, _ = parsear_fechas(r[_col(c, "dias del mes")], anio, mes)
                if isinstance(f_ini, datetime) and isinstance(f_fin, datetime):
                    fechas = [d for d in fechas if f_ini.date() <= d <= f_fin.date()] or fechas
                if not fechas and isinstance(f_ini, datetime):
                    fechas = [f_ini.date()]
                hi, hf = corregir_rango(parsear_hora(r[_col(c, "hora inicio")]), parsear_hora(r[_col(c, "hora fin")]))
                filas.append(Fila(titulo, _limpio(r[_col(c, "descripcion de la actividad")]), fechas, hi, hf,
                                  resolver_sede(_limpio(r[_col(c, "uva")]), sedes),
                                  _limpio(r[_col(c, "tipo de publico")]), _limpio(r[_col(c, "tipo de actividad")]),
                                  hoja=ws.title, fila=n))
        elif (enc := _encabezado(ws, ["titulo del curso", "fecha(s)"])):  # ── BIBLIOTECA
            res.formatos[ws.title] = "biblioteca"
            fila_h, c = enc
            titulo_hoja = " ".join(_limpio(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if x)
            mes, anio = _mes_de_titulo(titulo_hoja + " " + nombre_n)
            anio = anio or hoy.year
            for n, r in _filas(ws, fila_h + 1):
                nombre = r[_col(c, "titulo del curso")]
                if not nombre:
                    continue
                fechas, tent = parsear_fechas(r[_col(c, "fecha(s)")], anio, mes, str(r[_col(c, "dia(s)")] or ""))
                if not fechas:
                    res.descartados.append({"fila": f"{ws.title}#{n}", "titulo": _titulo(nombre), "motivo": "fecha_no_reconocida"})
                    continue
                horario = _norm(r[_col(c, "horario")]).split(" a ")
                hi = parsear_hora(horario[0]) if horario and horario[0] else None
                hf = parsear_hora(horario[1]) if len(horario) > 1 else None
                hi, hf = corregir_rango(hi, hf)
                insc = _norm(r[_col(c, "inscripcion")]).startswith("requiere")
                enlace = _limpio(r[_col(c, "enlace de inscripcion")])
                filas.append(Fila(_titulo(nombre), _limpio(r[_col(c, "descripcion")]), fechas, hi, hf,
                                  resolver_sede("biblioteca epm", sedes), _limpio(r[_col(c, "publico")]), "",
                                  _limpio(r[_col(c, "lugar")]), enlace if enlace.startswith("http") else "", insc,
                                  tent, hoja=ws.title, fila=n))
        else:
            continue
        for f in filas:
            evs = _a_eventos(f, hoy)
            if not evs:
                res.descartados.append({"fila": f"{f.hoja}#{f.fila}", "titulo": f.titulo, "motivo": "fechas_pasadas"})
            res.eventos.extend(evs)
    return res


def importar_programacion(data: bytes, nombre_archivo: str = "", aplicar: bool = True) -> dict:
    """Lee el Excel y pasa cada evento por la puerta de calidad (upsert por slug: idempotente)."""
    from collections import Counter
    res = leer_programacion(data, nombre_archivo)
    decisiones: Counter = Counter()
    sin_mapa = sorted({e["nombre_lugar"] for e in res.eventos if e.get("lat") is None and e.get("nombre_lugar")})
    if aplicar:
        from app.services.event_gate import insertar_evento
        for ev in res.eventos:
            payload = {k: v for k, v in ev.items() if not k.startswith("_")}
            try:
                decisiones[insertar_evento(payload, upsert=True).decision] += 1
            except Exception as exc:
                decisiones["error"] += 1
                print(f"[programacion_excel] {ev['titulo'][:50]}: {exc}")
    return {
        "archivo": nombre_archivo, "reparado": res.reparado, "formatos": res.formatos,
        "eventos": len(res.eventos), "decisiones": dict(decisiones),
        "descartados": Counter(d["motivo"] for d in res.descartados),
        "sedes_sin_coordenadas": sin_mapa,
        "muestra": [{k: e[k] for k in ("titulo", "fecha_inicio", "nombre_lugar")} for e in res.eventos[:8]],
    }


def importar_semillas(base: Optional[Path] = None) -> dict:
    """Importa una sola vez los Excel guardados en seeds/data/programacion/AAAA-MM/
    (marca por hash en config_kv). Se ejecuta tras cada deploy."""
    import hashlib
    from app.database import supabase
    base = base or SEDES_JSON.parent / "programacion"
    hoy = datetime.now(CO_TZ).date()
    resumen = {}
    for carpeta in sorted(p for p in base.glob("*") if p.is_dir()):
        try:
            y, m = map(int, carpeta.name.split("-"))
            if date(y, m, 28) < hoy - timedelta(days=31):
                continue  # meses ya pasados
        except ValueError:
            continue
        for f in sorted(carpeta.glob("*.xlsx")):
            data = f.read_bytes()
            clave = f"programacion_importada:{hashlib.sha1(data).hexdigest()[:16]}"
            try:
                if supabase.table("config_kv").select("key").eq("key", clave).execute().data:
                    continue
            except Exception:
                pass
            r = importar_programacion(data, f.name, aplicar=True)
            resumen[f"{carpeta.name}/{f.name}"] = {"eventos": r["eventos"], "decisiones": r["decisiones"]}
            try:
                supabase.table("config_kv").upsert({"key": clave, "value": json.dumps(resumen[f"{carpeta.name}/{f.name}"])},
                                                   on_conflict="key").execute()
                supabase.table("scraping_log").insert({
                    "fuente": "programacion_excel", "registros_nuevos": r["decisiones"].get("publicar", 0),
                    "registros_actualizados": r["decisiones"].get("duplicado", 0), "errores": r["decisiones"].get("error", 0),
                    "detalle": {"archivo": f"{carpeta.name}/{f.name}", **r["decisiones"],
                                "sin_coordenadas": r["sedes_sin_coordenadas"]},
                }).execute()
            except Exception as exc:
                print(f"[programacion_excel] marca/log: {exc}")
    print(f"📅 Programación Excel importada: {resumen or 'nada nuevo'}")
    return resumen
